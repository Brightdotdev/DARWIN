#!/usr/bin/env python3
"""
Test script for the document parser.
Run: python test_parser.py [path/to/pdf1.pdf] [path/to/pdf2.pdf]
"""

import sys
import json
from pathlib import Path
from datetime import datetime

# Add src to path if needed
sys.path.insert(0, str(Path(__file__).parent))

from src.core.Parsing.Parser import parse, parse_many, scan_directory
from src.core.Parsing.ParserUtils import ParsedDocument


# ── Helper to create test files ─────────────────────────────────────────────

def create_test_docx(path: Path, content: str = None):
    """Create a test DOCX file programmatically."""
    try:
        from docx import Document
        from docx.shared import Inches
        
        doc = Document()
        doc.add_heading('Test Document', 0)
        
        if content:
            doc.add_paragraph(content)
        else:
            doc.add_paragraph('This is a test DOCX file created for parser testing.')
            doc.add_paragraph('It contains multiple paragraphs to test extraction.')
            doc.add_heading('Section 1', level=1)
            doc.add_paragraph('First section content with **bold** and *italic* text.')
            doc.add_heading('Section 2', level=1)
            doc.add_paragraph('Second section with a bullet list:', style='List Bullet')
            doc.add_paragraph('Item one', style='List Bullet')
            doc.add_paragraph('Item two', style='List Bullet')
            doc.add_paragraph('Item three', style='List Bullet')
        
        doc.save(str(path))
        return True
    except ImportError:
        print("⚠️  python-docx not installed - run: pip install python-docx")
        return False


def create_test_pptx(path: Path, content: str = None):
    """Create a test PPTX file programmatically."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
        
        prs = Presentation()
        
        # Title slide
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = title_slide.shapes.title
        subtitle = title_slide.placeholders[1]
        title.text = "Test Presentation"
        subtitle.text = "Created for parser testing"
        
        # Content slide
        bullet_slide = prs.slides.add_slide(prs.slide_layouts[1])
        shapes = bullet_slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = "Key Points"
        tf = body_shape.text_frame
        tf.text = "First bullet point"
        
        p = tf.add_paragraph()
        p.text = "Second bullet point"
        p.level = 1
        
        p = tf.add_paragraph()
        p.text = "Third bullet point"
        p.level = 2
        
        if content:
            # Add another slide with custom content
            custom_slide = prs.slides.add_slide(prs.slide_layouts[1])
            custom_slide.shapes.title.text = "Custom Content"
            custom_slide.placeholders[1].text = content
        
        prs.save(str(path))
        return True
    except ImportError:
        print("⚠️  python-pptx not installed - run: pip install python-pptx")
        return False


# ── Test functions ──────────────────────────────────────────────────────────

def print_doc_summary(doc: ParsedDocument, index: int = None):
    """Print a formatted summary of a parsed document."""
    prefix = f"[{index}] " if index is not None else ""
    
    print(f"\n{'='*70}")
    print(f"{prefix}{doc.filename}")
    print(f"{'='*70}")
    
    if doc.error:
        print(f"❌ ERROR: {doc.error}")
        return
    
    print(f"✅ Type: {doc.file_type} | Strategy: {doc.strategy}")
    print(f"📄 Pages: {doc.page_count} | Size: {len(doc.text):,} chars")
    
    if doc.images:
        print(f"🖼️  Images: {len(doc.images)}")
        for img in doc.images[:3]:  # Show first 3
            print(f"   - Page {img.page}: {img.path.name} ({img.width}x{img.height})")
            if img.caption:
                print(f"     Caption: {img.caption[:80]}...")
        if len(doc.images) > 3:
            print(f"   ... and {len(doc.images) - 3} more")
    
    # Show first 500 chars of text
    print(f"\n📝 Preview (first 500 chars):")
    print("-" * 50)
    print(doc.text[:500])
    if len(doc.text) > 500:
        print(f"\n... (truncated, {len(doc.text) - 500} more chars)")
    
    return doc.text


def test_single_file():
    """Test parsing a single file."""
    print("\n" + "🧪 TEST 1: Single File Parsing".center(70))
    
    # Create a test text file
    test_file = Path("test_sample.txt")
    test_file.write_text("""# Test Document

This is a test document for the parser.

## Section 1

The parser should extract this text and convert it to markdown.

## Section 2

It should handle multiple paragraphs and basic markdown formatting.

- Bullet point 1
- Bullet point 2
- Bullet point 3

That's it for the test.
""")
    
    doc = parse(test_file)
    print_doc_summary(doc)
    
    # Cleanup
    test_file.unlink()
    return doc


def test_docx():
    """Test parsing a DOCX file."""
    print("\n" + "🧪 TEST 2: DOCX Parsing".center(70))
    
    test_file = Path("test_sample.docx")
    if create_test_docx(test_file):
        doc = parse(test_file)
        print_doc_summary(doc)
        test_file.unlink()
        return doc
    
    print("❌ Could not create DOCX file")
    return None


def test_pptx():
    """Test parsing a PPTX file."""
    print("\n" + "🧪 TEST 3: PPTX Parsing".center(70))
    
    test_file = Path("test_sample.pptx")
    if create_test_pptx(test_file):
        doc = parse(test_file)
        print_doc_summary(doc)
        test_file.unlink()
        return doc
    
    print("❌ Could not create PPTX file")
    return None


def test_specific_pdfs(pdf_paths: list[Path]):
    """Test parsing user-specified PDF files."""
    print("\n" + "🧪 TEST 4: User-Specified PDFs".center(70))
    
    if not pdf_paths:
        print("⚠️  No PDF paths provided")
        return []
    
    results = []
    for pdf_path in pdf_paths:
        if pdf_path.exists() and pdf_path.suffix.lower() == '.pdf':
            print(f"\n📄 Testing: {pdf_path.name}")
            doc = parse(pdf_path)
            print_doc_summary(doc)
            results.append(doc)
        else:
            print(f"❌ Invalid PDF: {pdf_path}")
    
    return results


def test_multiple_files():
    """Test parsing multiple files at once."""
    print("\n" + "🧪 TEST 5: Multiple Files".center(70))
    
    # Create test files
    files = []
    
    # TXT file
    txt_file = Path("test1.txt")
    txt_file.write_text("This is test file 1. It contains simple text.")
    files.append(txt_file)
    
    # MD file
    md_file = Path("test2.md")
    md_file.write_text("# Markdown Test\n\nThis is **bold** and *italic* text.\n\n- List item 1\n- List item 2")
    files.append(md_file)
    
    # Another TXT
    txt2_file = Path("test3.txt")
    txt2_file.write_text("Test file 3 with different content to parse.")
    files.append(txt2_file)
    
    # Create DOCX if available
    docx_file = Path("test4.docx")
    if create_test_docx(docx_file, "Additional test document"):
        files.append(docx_file)
    
    # Parse all
    docs = parse_many(files)
    
    for i, doc in enumerate(docs, 1):
        print_doc_summary(doc, i)
    
    # Cleanup
    for f in files:
        if f.exists():
            f.unlink()
    
    return docs


def test_with_images():
    """Test parsing with image extraction (requires a PDF with images)."""
    print("\n" + "🧪 TEST 6: Image Extraction".center(70))
    print("⚠️  This test requires a PDF with images.")
    print("   Place PDFs in the current directory or pass as arguments")
    
    # Look for any PDF with images
    pdf_files = list(Path(".").glob("*.pdf"))
    pdf_with_images = None
    
    for pdf in pdf_files:
        if pdf.stat().st_size > 10000:  # Assume >10KB might have images
            pdf_with_images = pdf
            break
    
    if not pdf_with_images:
        print("❌ No PDF found for image test - skipping")
        return None
    
    images_dir = Path("test_images_output")
    doc = parse(pdf_with_images, images_dir=images_dir, caption=False)
    print_doc_summary(doc)
    
    if images_dir.exists() and doc.images:
        print(f"\n📁 Images saved to: {images_dir}")
        for img in doc.images[:5]:
            print(f"   - {img.path.name} ({img.width}x{img.height})")
    
    return doc


def test_directory_scan():
    """Test scanning an entire directory."""
    print("\n" + "🧪 TEST 7: Directory Scan".center(70))
    
    # Create a test directory
    test_dir = Path("test_scan_dir")
    test_dir.mkdir(exist_ok=True)
    
    # Create various files
    (test_dir / "doc1.txt").write_text("First document in directory")
    (test_dir / "doc2.md").write_text("# Second document\n\nContent here")
    (test_dir / "doc3.txt").write_text("Third document with more text")
    (test_dir / "ignore.exe").write_text("binary fake")  # Should be ignored
    (test_dir / ".hidden.txt").write_text("hidden")  # Should be ignored
    
    # Create DOCX if available
    create_test_docx(test_dir / "doc4.docx", "Document in subfolder")
    
    sub_dir = test_dir / "subfolder"
    sub_dir.mkdir()
    (sub_dir / "doc5.txt").write_text("Nested document")
    
    # Test non-recursive
    print("\n📂 Non-recursive scan:")
    docs = scan_directory(test_dir, recursive=False)
    print(f"   Found {len(docs)} files")
    for doc in docs:
        print(f"   - {doc.filename}")
    
    # Test recursive
    print("\n📂 Recursive scan:")
    docs = scan_directory(test_dir, recursive=True)
    print(f"   Found {len(docs)} files")
    for doc in docs:
        print(f"   - {doc.filename}")
        if doc.error:
            print(f"     ❌ {doc.error[:100]}")
        else:
            print(f"     ✅ {len(doc.text):,} chars, {doc.page_count} pages")
    
    # Cleanup
    import shutil
    shutil.rmtree(test_dir)
    
    return docs


def test_error_handling():
    """Test parsing of invalid files."""
    print("\n" + "🧪 TEST 8: Error Handling".center(70))
    
    test_cases = [
        ("nonexistent.pdf", "Missing file"),
        ("test.xyz", "Unsupported type"),
    ]
    
    for filename, description in test_cases:
        print(f"\n📄 {description}: {filename}")
        doc = parse(filename)
        if doc.error:
            print(f"   ✅ Caught error: {doc.error[:80]}...")
        else:
            print(f"   ⚠️  No error returned (unexpected)")
    
    # Create a corrupted file
    corrupt_file = Path("corrupt.txt")
    corrupt_file.write_bytes(b"\x00\x01\x02\x03corrupt data")
    doc = parse(corrupt_file)
    print(f"\n📄 Corrupted file test:")
    if doc.error:
        print(f"   ✅ Caught error: {doc.error[:80]}")
    else:
        print(f"   ⚠️  No error (text: {doc.text[:50]})")
    corrupt_file.unlink()


def test_performance():
    """Test parsing performance with a larger file."""
    print("\n" + "🧪 TEST 9: Performance".center(70))
    
    # Create a larger text file
    large_file = Path("large_test.txt")
    content = "\n".join([f"Line {i}: This is test content for performance testing." 
                         for i in range(1000)])
    large_file.write_text(content)
    
    start = datetime.now()
    doc = parse(large_file)
    elapsed = (datetime.now() - start).total_seconds()
    
    print(f"📄 File size: {large_file.stat().st_size:,} bytes")
    print(f"⏱️  Parse time: {elapsed:.2f} seconds")
    print(f"📝 Extracted: {len(doc.text):,} characters")
    print(f"⚡ Speed: {len(doc.text)/elapsed:,.0f} chars/second")
    
    large_file.unlink()
    return doc


def test_output_format():
    """Demonstrate the exact output format."""
    print("\n" + "🧪 TEST 10: Output Format".center(70))
    
    # Create a test file
    test_file = Path("output_test.md")
    test_file.write_text("""# Output Format Test

This document demonstrates what the parser returns.

## What you get:

1. **ParsedDocument object** with fields:
   - `.text` - Clean markdown text
   - `.filename` - Original filename
   - `.page_count` - Number of pages
   - `.file_type` - 'pdf' or 'markitdown'
   - `.metadata` - Document metadata
   - `.images` - List of extracted images
   - `.error` - None if success

2. The `.text` field contains markdown ready for chunking.
""")
    
    doc = parse(test_file)
    
    print("\n📦 ParsedDocument object:")
    print(f"   Type: {type(doc).__name__}")
    print(f"   Fields: {[f for f in dir(doc) if not f.startswith('_')]}")
    
    print("\n📝 .text content (markdown):")
    print("-" * 50)
    print(doc.text)
    print("-" * 50)
    
    print("\n💾 You can save the markdown to a file:")
    print(f"   with open('output.md', 'w') as f: f.write(doc.text)")
    
    test_file.unlink()
    return doc


def run_all_tests(pdf_paths: list[Path] = None):
    """Run all tests and generate a report."""
    print("\n" + "🚀 DOCUMENT PARSER TEST SUITE".center(70))
    print(f"Started at: {datetime.now()}")
    print(f"Python: {sys.version}")
    
    results = {
        "single_file": test_single_file(),
        "docx": test_docx(),
        "pptx": test_pptx(),
        "specific_pdfs": test_specific_pdfs(pdf_paths) if pdf_paths else None,
        "multiple_files": test_multiple_files(),
        "with_images": test_with_images(),
        "directory_scan": test_directory_scan(),
        "error_handling": test_error_handling(),
        "performance": test_performance(),
        "output_format": test_output_format(),
    }
    
    # Summary
    print("\n" + "📊 TEST SUMMARY".center(70))
    print(f"{'Test':<20} {'Status':<10} {'Pages':<8} {'Chars':<12}")
    print("-" * 70)
    
    for name, doc in results.items():
        if doc is None:
            status = "⏭️ SKIP"
            pages = "-"
            chars = "-"
        elif isinstance(doc, list):
            # Handle list of PDF results
            if doc:
                status = f"✅ {len(doc)} PDFs"
                pages = sum(d.page_count for d in doc if not d.error)
                chars = sum(len(d.text) for d in doc if not d.error)
            else:
                status = "⏭️ SKIP"
                pages = "-"
                chars = "-"
        elif doc.error:
            status = "❌ FAIL"
            pages = "-"
            chars = "-"
        else:
            status = "✅ PASS"
            pages = doc.page_count
            chars = f"{len(doc.text):,}"
        
        print(f"{name:<20} {status:<10} {pages:<8} {chars:<12}")
    
    print("\n✨ Test suite complete!")
    print("\n💡 Expected Output: ParsedDocument objects with clean markdown in .text field")
    print("   Use doc.text for chunking, doc.metadata for citations, doc.images for vision")


if __name__ == "__main__":
    # Parse command line arguments for PDF files
    pdf_paths = []
    for arg in sys.argv[1:]:
        pdf_path = Path(arg)
        if pdf_path.exists() and pdf_path.suffix.lower() == '.pdf':
            pdf_paths.append(pdf_path)
        else:
            print(f"⚠️  Warning: '{arg}' is not a valid PDF file, skipping...")
    
    if pdf_paths:
        print(f"📚 Found {len(pdf_paths)} PDF(s) to test:")
        for p in pdf_paths:
            print(f"   - {p.name}")
    
    run_all_tests(pdf_paths)