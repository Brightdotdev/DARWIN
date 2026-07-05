import re
import csv
import email
import zipfile
from pathlib import Path
from typing import Optional, Tuple, Dict, List

from src.core.Pipeline.Ingestion.Parsing.lib import ExtractedImage


def _words_to_pages(text: str, words_per_page: int = 275) -> int:
    """Very rough page count for non‑PDF documents."""
    word_count = len(text.split())
    return max(1, round(word_count / words_per_page))


# ═══════════════════════════════════════════════════════════════════════════
# HTML / XML
# ═══════════════════════════════════════════════════════════════════════════

def parse_html(
    path: Path,
    images_dir: Optional[Path] = None,
    caption: bool = False,
    precise: bool = False,
) -> Tuple[str, str, int, Dict, List[ExtractedImage]]:
    """Extract readable text from an HTML file using Trafilatura."""
    try:
        from trafilatura import extract
    except ImportError:
        raise ImportError("trafilatura is required for HTML parsing. Install with: pip install trafilatura")

    html = path.read_text(encoding="utf-8", errors="replace")

    text = extract(
        html,
        output_format="txt",
        favor_recall=precise,
        include_comments=False,
        include_tables=True,
        deduplicate=True,
    )

    if text is None:
        text = ""

    page_count = _words_to_pages(text)
    return text, text, page_count, {}, []


def parse_xml(
    path: Path,
    images_dir: Optional[Path] = None,
    caption: bool = False,
    precise: bool = False,
) -> Tuple[str, str, int, Dict, List[ExtractedImage]]:
    """Extract text from an XML file (generic)."""
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        raise ImportError("beautifulsoup4 is required. Install with: pip install beautifulsoup4")

    xml_content = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(xml_content, "xml")
    text = soup.get_text(separator="\n", strip=True)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    page_count = _words_to_pages(text)
    return text, text, page_count, {}, []


# ═══════════════════════════════════════════════════════════════════════════
# CSV / TSV
# ═══════════════════════════════════════════════════════════════════════════

def parse_csv(
    path: Path,
    images_dir: Optional[Path] = None,
    caption: bool = False,
    precise: bool = False,
) -> Tuple[str, str, int, Dict, List[ExtractedImage]]:
    """Convert CSV/TSV to a Markdown table."""
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","

    with open(path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        rows = list(reader)

    if not rows:
        return "", "", 0, {}, []

    header = rows[0]
    header_cells = [cell.replace("|", "\\|") for cell in header]
    md = "| " + " | ".join(header_cells) + " |\n"
    md += "| " + " | ".join(["---"] * len(header)) + " |\n"
    for row in rows[1:]:
        cells = [cell.replace("|", "\\|") for cell in row]
        while len(cells) < len(header):
            cells.append("")
        md += "| " + " | ".join(cells[:len(header)]) + " |\n"

    page_count = _words_to_pages(md)
    return md, md, page_count, {}, []


# ═══════════════════════════════════════════════════════════════════════════
# Spreadsheets (.xlsx, .xls, .ods)
# ═══════════════════════════════════════════════════════════════════════════

def parse_spreadsheet(
    path: Path,
    images_dir: Optional[Path] = None,
    caption: bool = False,
    precise: bool = False,
) -> Tuple[str, str, int, Dict, List[ExtractedImage]]:
    """Convert a spreadsheet file to Markdown tables (one per sheet)."""
    ext = path.suffix.lower()

    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl is required for Excel files. Install with: pip install openpyxl")
        wb = openpyxl.load_workbook(path, data_only=True)
    elif ext == ".ods":
        try:
            import pyexcel_ods
        except ImportError:
            raise ImportError("pyexcel-ods is required for ODS files. Install with: pip install pyexcel-ods")
        data = pyexcel_ods.get_data(str(path))
        return _parse_ods(path, images_dir, caption, precise)
    else:
        raise ValueError(f"Unsupported spreadsheet extension: {ext}")

    md_parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        md_parts.append(f"## {sheet_name}\n")
        header = [str(cell) if cell is not None else "" for cell in rows[0]]
        md_parts.append("| " + " | ".join(header) + " |")
        md_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
        for row in rows[1:]:
            cells = [str(cell).replace("|", "\\|") if cell is not None else "" for cell in row]
            cells = cells[:len(header)]
            while len(cells) < len(header):
                cells.append("")
            md_parts.append("| " + " | ".join(cells) + " |")
        md_parts.append("")

    full_md = "\n".join(md_parts).strip()
    page_count = _words_to_pages(full_md)
    return full_md, full_md, page_count, {}, []


def _parse_ods(path: Path, images_dir, caption, precise):
    import pyexcel_ods
    data = pyexcel_ods.get_data(str(path))
    md_parts = []
    for sheet_name, rows in data.items():
        if not rows:
            continue
        md_parts.append(f"## {sheet_name}\n")
        header = [str(cell) if cell is not None else "" for cell in rows[0]]
        md_parts.append("| " + " | ".join(header) + " |")
        md_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
        for row in rows[1:]:
            cells = [str(cell).replace("|", "\\|") if cell is not None else "" for cell in row]
            cells = cells[:len(header)]
            while len(cells) < len(header):
                cells.append("")
            md_parts.append("| " + " | ".join(cells) + " |")
        md_parts.append("")
    full_md = "\n".join(md_parts).strip()
    page_count = _words_to_pages(full_md)
    return full_md, full_md, page_count, {}, []


# ═══════════════════════════════════════════════════════════════════════════
# E‑book (.epub)
# ═══════════════════════════════════════════════════════════════════════════

def parse_ebook(
    path: Path,
    images_dir: Optional[Path] = None,
    caption: bool = False,
    precise: bool = False,
) -> Tuple[str, str, int, Dict, List[ExtractedImage]]:
    """Extract text and optionally images from an EPUB file."""
    from bs4 import BeautifulSoup
    import os

    with zipfile.ZipFile(path, "r") as zf:
        xhtml_files = sorted([
            name for name in zf.namelist()
            if name.endswith((".xhtml", ".html", ".htm"))
            and not name.startswith("__")
        ])

        all_text = []
        all_images = []

        for name in xhtml_files:
            content = zf.read(name).decode("utf-8", errors="replace")
            soup = BeautifulSoup(content, "html.parser")

            if images_dir:
                for img_tag in soup.find_all("img"):
                    src = img_tag.get("src", "")
                    if not src:
                        continue
                    img_zip_path = str(Path(name).parent / src) if not src.startswith("/") else src[1:]
                    try:
                        img_data = zf.read(img_zip_path)
                        img_name = Path(src).name
                        img_path = images_dir / img_name
                        img_path.write_bytes(img_data)
                        all_images.append({
                            "path": str(img_path),
                            "page": len(all_text),
                            "index": len(all_images),
                            "caption": "",
                            "width": 0,
                            "height": 0,
                            "slot": len(all_images),
                        })
                    except (KeyError, FileNotFoundError):
                        pass

            for img_tag in soup.find_all("img"):
                img_tag.decompose()

            text = soup.get_text(separator="\n", strip=True)
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            all_text.append(text)

        full_text = "\n\n---\n\n".join(all_text)
        page_count = _words_to_pages(full_text)

        extracted = [ExtractedImage(
            path=Path(d["path"]),
            page=d["page"],
            index=d["index"],
            caption=d["caption"],
            width=d.get("width", 0),
            height=d.get("height", 0)
        ) for d in all_images]

        return full_text, full_text, page_count, {}, extracted


# ═══════════════════════════════════════════════════════════════════════════
# Email (.eml)
# ═══════════════════════════════════════════════════════════════════════════

def parse_email(
    path: Path,
    images_dir: Optional[Path] = None,
    caption: bool = False,
    precise: bool = False,
) -> Tuple[str, str, int, Dict, List[ExtractedImage]]:
    """Extract text from an email file (.eml)."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        msg = email.message_from_file(f)

    headers = {
        "Subject": msg.get("Subject", ""),
        "From": msg.get("From", ""),
        "To": msg.get("To", ""),
        "Date": msg.get("Date", ""),
    }

    body_parts = []
    if msg.is_multipart():
        for part in msg.walk():
            content_disposition = str(part.get("Content-Disposition", ""))
            if "attachment" in content_disposition:
                continue
            payload = part.get_payload(decode=True)
            if payload:
                charset = part.get_content_charset() or "utf-8"
                try:
                    text = payload.decode(charset, errors="replace")
                except (LookupError, UnicodeDecodeError):
                    text = payload.decode("utf-8", errors="replace")
                if part.get_content_type() == "text/html":
                    try:
                        from bs4 import BeautifulSoup
                        soup = BeautifulSoup(text, "html.parser")
                        text = soup.get_text(separator="\n", strip=True)
                    except ImportError:
                        pass
                body_parts.append(text)
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            charset = msg.get_content_charset() or "utf-8"
            try:
                text = payload.decode(charset, errors="replace")
            except (LookupError, UnicodeDecodeError):
                text = payload.decode("utf-8", errors="replace")
            if msg.get_content_type() == "text/html":
                try:
                    from bs4 import BeautifulSoup
                    soup = BeautifulSoup(text, "html.parser")
                    text = soup.get_text(separator="\n", strip=True)
                except ImportError:
                    pass
            body_parts.append(text)

    body = "\n".join(body_parts).strip()

    lines = []
    if headers["Subject"]:
        lines.append(f"Subject: {headers['Subject']}")
    if headers["From"]:
        lines.append(f"From: {headers['From']}")
    if headers["To"]:
        lines.append(f"To: {headers['To']}")
    if headers["Date"]:
        lines.append(f"Date: {headers['Date']}")
    lines.append("")
    lines.append(body)

    full_text = "\n".join(lines).strip()
    page_count = _words_to_pages(full_text)
    return full_text, full_text, page_count, {}, []